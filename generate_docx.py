from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

path = r"D:\Claude_VSCode_20260618\KCG\KoreaStrategyGroup\정관_코리아전략그룹.docx"

doc = Document()
# set default font to Malgun Gothic for Korean compatibility
style = doc.styles['Normal']
font = style.font
font.name = 'Malgun Gothic'
try:
    font.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
except Exception:
    pass

# Title
p = doc.add_paragraph()
p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
r = p.add_run('정관')
r.bold = True
r.font.size = Pt(18)
try:
    r.font.name = 'Malgun Gothic'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
except Exception:
    pass

p2 = doc.add_paragraph()
p2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
r2 = p2.add_run('(주식회사 코리아전략그룹)')
r2.font.size = Pt(12)

# Helper to add article
def article(title, content_lines):
    h = doc.add_paragraph()
    run = h.add_run(title)
    run.bold = True
    run.font.size = Pt(12)
    for line in content_lines:
        doc.add_paragraph(line)

# Articles
article('제1조 (상호)', ['본 회사는 주식회사 코리아전략그룹(이하 "회사"라 한다)이라 칭한다.'])

article('제2조 (목적)', [
    '회사는 다음의 사업을 영위함을 목적으로 한다.',
    '1. 공공 및 민간 부문에 대한 정책연구, 여론조사 및 데이터 분석',
    '2. 공공사업 및 정책사업 기획·관리·평가 용역',
    '3. 정치컨설팅, 선거전략 및 여론관리 서비스',
    '4. 디지털 캠페인·홍보 전략 및 미디어 집행',
    '5. 교육·세미나·컨퍼런스 기획 및 운영',
    '6. 연구보고서·백서 작성 및 출판',
    '7. 위 각호에 부대되는 일체의 사업'
])

article('제3조 (본점의 소재지)', ['회사의 본점은 대한민국에 둔다. 필요에 따라 이사회 결의로 국내외에 지점이나 사무소를 둘 수 있다.'])

article('제4조 (관할 법원)', ['회사의 소송에 관하여는 본점 소재지의 관할법원을 관할법원으로 한다.'])

article('제5조 (회사의 존속기간)', ['회사의 존속기간은 정함이 없으며, 정관의 변경으로 정할 수 있다.'])

article('제6조 (자본금)', ['회사의 자본금은 금 삼억원정(￦300,000,000)으로 하고, 액면가액은 1주당 금 일십만원정(￦100,000)으로 하여 총 3,000주를 발행한다. 자본금 및 발행주식수는 주주총회의 결의로 변경할 수 있다.'])

article('제7조 (주식의 양도)', ['주식의 양도에 관하여는 상법 및 본 정관의 규정을 따른다. 필요시 이사회 또는 주주총회의 동의를 요할 수 있다.'])

article('제8조 (주주총회)', ['주주총회는 정기주주총회와 임시주주총회로 구분하며, 정기주주총회는 매 사업연도 종료 후 정해진 시기에 본 회사의 결산을 승인하고 기타 법령 또는 정관에서 정하는 사항을 심의·의결한다.'])

article('제9조 (이사회)', [
    '1. 회사는 이사회를 둔다.',
    '2. 이사회의 구성은 이사 3인 이상으로 하며, 이사의 선임·해임 및 이사회 운영에 관한 사항은 상법 및 내부규정에 따른다.',
    '3. 이사회는 회사의 주요 경영정책·사업계획의 승인, 예산·결산의 승인, 대표이사의 선임 및 해임 등을 의결한다.'
])

article('제10조 (대표이사)', [
    '1. 회사는 대표이사를 둔다.',
    '2. 대표이사는 정도현(鄭度賢)을 초대 대표이사로 선임한다.',
    '3. 대표이사는 회사의 업무를 집행하고, 이사회의 결의사항을 집행한다.'
])

article('제11조 (임원)', ['1. 회사는 필요에 따라 감사 및 기타 임원을 둘 수 있다.', '2. 임원의 권한과 의무는 관련 법령 및 이사회 결의에 따른다.'])

article('제12조 (사업부문)', [
    '회사는 다음의 3개 부문을 운영한다.',
    '1. 연구부문(연구개발 및 정책연구 담당)',
    '2. 과제수행부문(공공·민간 용역의 기획·관리·실행 담당)',
    '3. 정치컨설팅부문(정치 전략, 여론조사 및 캠페인 담당)',
    '각 부문은 대표이사의 지휘·감독을 받으며, 세부 직제 및 운영은 내부규정에 따른다.'
])

article('제13조 (회계연도 및 결산)', ['회사의 회계연도는 매년 1월 1일부터 12월 31일까지로 한다. 매 결산기에 대하여 회계감사 등을 거쳐 손익계산서 및 재무제표를 작성하여 주주총회의 승인을 받는다.'])

article('제14조 (정관 변경)', ['정관의 변경은 주주총회에서 법령이 정하는 특별결의를 거쳐야 한다.'])

article('제15조 (해산 및 청산)', ['회사는 법령 또는 주주총회의 결의에 의해 해산할 수 있으며, 해산 시에는 청산절차를 거친다.'])

article('제16조 (기타)', ['이 정관에 규정되지 않은 사항은 상법 및 기타 관련 법령에 따른다.'])

# 부칙
p = doc.add_paragraph()
p.add_run('부칙').bold = True
p = doc.add_paragraph()
p.add_run('이 정관은 주주총회의 결의로서 확정된 날로부터 시행한다.')

# Save
try:
    doc.save(path)
    print('DOCX_GENERATED')
except Exception as e:
    print('DOCX_SAVE_FAILED:', e)
